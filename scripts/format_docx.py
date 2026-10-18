# /// script
# requires-python = ">=3.10"
# dependencies = [
#   "python-docx>=1.1",
# ]
# ///
"""
Post-process pandoc-generated docx:
  - adds thin gray borders to all tables
  - sets column widths per table using TABLE_WIDTHS_BY_STEM, keyed on the input
    filename stem (basename without .docx), listing one row of column fractions
    per table in the order they appear in the document

Usage:
    uv run scripts/format_docx.py <input.docx>                 # overwrite in place
    uv run scripts/format_docx.py <input.docx> <output.docx>   # separate output
    uv run scripts/format_docx.py <input.docx> --borders-only  # skip column widths

The file is overwritten in-place if no output path is given.

Adapted from the same-named script in ../fjordgata30/scripts/. Mechanics are
identical; the difference is that this version is keyed on filename stem via
TABLE_WIDTHS_BY_STEM rather than a positional list, so one script can serve
multiple deliverables without the "one new table in doc A shifts doc B's widths
by one" problem.
"""

import sys
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Text width on an A4 page with 2.5 cm margins on each side: 21 - 5 = 16 cm.
# Column widths are expressed as fractions of this total.
PAGE_WIDTH_CM = 16.0

BORDER_COLOR = "BFBFBF"   # light gray
BORDER_SIZE  = 4           # in eighths of a point (4 = 0.5 pt)

# Column widths (fractions of PAGE_WIDTH_CM), per input filename stem, listed
# in the order tables appear in the document. A row of fractions must sum to
# <= 1.0 (Pandoc/Word ignores the remainder). Add a new entry when a new
# deliverable with tables is produced.
TABLE_WIDTHS_BY_STEM: dict[str, list[list[float]]] = {
    # Example (uncomment and adapt when adding a new stem):
    # "2026-10-16_my_deliverable": [
    #     [0.50, 0.50],
    #     [0.33, 0.33, 0.34],
    # ],
}


def _get_or_add(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el


def _set_border_element(borders_el, edge):
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"),   "single")
    el.set(qn("w:sz"),    str(BORDER_SIZE))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), BORDER_COLOR)
    borders_el.append(el)


def add_cell_borders(cell):
    tc = cell._tc
    tcPr = _get_or_add(tc, "w:tcPr")
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _set_border_element(tcBorders, edge)
    tcPr.append(tcBorders)


def set_column_widths(table, fractions):
    """Set column widths (in twips) for every cell in the table."""
    widths_twips = [int(f * PAGE_WIDTH_CM * 567) for f in fractions]
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i >= len(widths_twips):
                break
            tc = cell._tc
            tcPr = _get_or_add(tc, "w:tcPr")
            for old in tcPr.findall(qn("w:tcW")):
                tcPr.remove(old)
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"),    str(widths_twips[i]))
            tcW.set(qn("w:type"), "dxa")
            tcPr.insert(0, tcW)


def process(input_path, output_path, borders_only=False):
    doc = Document(input_path)
    stem = Path(input_path).stem
    widths = TABLE_WIDTHS_BY_STEM.get(stem)

    for i, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                add_cell_borders(cell)

        if borders_only:
            continue
        if widths is None:
            if i == 0:
                print(f"  note: no TABLE_WIDTHS_BY_STEM entry for '{stem}' — borders applied, column widths left as pandoc produced them. Add an entry if the layout needs tuning.")
            continue
        if i < len(widths):
            set_column_widths(table, widths[i])
        else:
            print(f"  warning: '{stem}' has {len(widths)} width rows but doc has {i + 1}+ tables — widths left as-is for table {i}")

    doc.save(output_path)
    print(f"Done: {len(doc.tables)} tables processed -> {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    borders_only = "--borders-only" in sys.argv
    inp = args[0]
    outp = args[1] if len(args) > 1 else inp
    process(inp, outp, borders_only=borders_only)
