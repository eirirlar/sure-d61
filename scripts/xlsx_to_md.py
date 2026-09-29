#!/usr/bin/env python
"""Convert an .xlsx workbook to a Markdown document with one table per sheet.

Usage:
    python xlsx_to_md.py <input.xlsx> <output.md>
"""

import sys
from pathlib import Path

from openpyxl import load_workbook


def cell_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("|", "\\|").replace("\n", " ").strip()


def sheet_to_md(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return f"### {ws.title}\n\n_(empty)_\n"

    max_cols = max(len(r) for r in rows)
    rows = [list(r) + [None] * (max_cols - len(r)) for r in rows]

    # Trim trailing all-None columns
    while max_cols > 0 and all(r[max_cols - 1] is None for r in rows):
        for r in rows:
            r.pop()
        max_cols -= 1
    # Trim trailing all-None rows
    while rows and all(c is None for c in rows[-1]):
        rows.pop()

    if not rows or max_cols == 0:
        return f"### {ws.title}\n\n_(empty)_\n"

    header = [cell_str(v) for v in rows[0]]
    body = rows[1:]

    out = [f"### {ws.title}", ""]
    out.append("| " + " | ".join(header) + " |")
    out.append("|" + "|".join("---" for _ in header) + "|")
    for r in body:
        out.append("| " + " | ".join(cell_str(v) for v in r) + " |")
    out.append("")
    return "\n".join(out)


def main():
    if len(sys.argv) != 3:
        print(__doc__, file=sys.stderr)
        sys.exit(1)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])

    wb = load_workbook(src, data_only=True)
    parts = [f"# {src.stem}", ""]
    for name in wb.sheetnames:
        parts.append(sheet_to_md(wb[name]))
    dst.write_text("\n".join(parts), encoding="utf-8")
    print(f"wrote {dst} ({len(wb.sheetnames)} sheet(s))")


if __name__ == "__main__":
    main()
